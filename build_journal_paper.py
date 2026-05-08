"""Build journal-paper.docx from journal_paper_publishing_blueprint.md.

Target: 10-page Indian Economic Journal style draft, around 5000 words,
APA author-year citations, British English, ADL-only, layered attenuation,
plain MS Economics student prose, no em dashes. All variables and equations
use real Word subscript and superscript runs (no literal x_t in body).
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

ROOT = str(Path(__file__).resolve().parent)
OUT = os.path.join(ROOT, "journal-paper.docx")

doc = Document()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    run = paragraph.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_page_number(section.footer.paragraphs[0])


def force_tnr(style):
    style.font.name = 'Times New Roman'
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(k), 'Times New Roman')


normal = doc.styles['Normal']
normal.font.size = Pt(12)
force_tnr(normal)
pf = normal.paragraph_format
pf.line_spacing = 1.15
pf.space_before = Pt(3)
pf.space_after = Pt(3)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.first_line_indent = Inches(0.25)

for name, sz, sp_b, sp_a in [('Heading 1', 13, 12, 6), ('Heading 2', 12, 8, 4)]:
    s = doc.styles[name]
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    force_tnr(s)
    s.paragraph_format.space_before = Pt(sp_b)
    s.paragraph_format.space_after = Pt(sp_a)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing = 1.15
    s.paragraph_format.first_line_indent = Inches(0)
doc.styles['Heading 2'].font.italic = True

cs = doc.styles['Caption']
cs.font.size = Pt(10)
cs.font.color.rgb = RGBColor(0, 0, 0)
force_tnr(cs)
cs.paragraph_format.line_spacing = 1.0


# -------- Inline math paragraph (sub/sup parsing) --------

def add_para_math(text, indent=True, justify=True, size=12, italic_default=False):
    """Write a paragraph that turns _x, _{xx}, ^x, ^{xx} into real
    subscript and superscript runs. Greek letters and Δ should be put as
    Unicode in the source string. The rest is plain Times New Roman."""
    p = doc.add_paragraph()
    if not indent:
        p.paragraph_format.first_line_indent = Inches(0)
    if not justify:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if ch in ('_', '^'):
            subscript = ch == '_'
            i += 1
            if i < n and text[i] == '{':
                j = text.find('}', i)
                if j == -1:
                    content = text[i:]
                    i = n
                else:
                    content = text[i + 1:j]
                    i = j + 1
            elif i < n:
                content = text[i]
                i += 1
            else:
                content = ''
            r = p.add_run(content)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(size)
            r.font.subscript = subscript
            r.font.superscript = not subscript
            if italic_default:
                r.italic = True
        else:
            j = i
            while j < n and text[j] not in ('_', '^'):
                j += 1
            r = p.add_run(text[i:j])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(size)
            if italic_default:
                r.italic = True
            i = j
    return p


def add_para(text, indent=True, justify=True, size=12, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    if not indent:
        p.paragraph_format.first_line_indent = Inches(0)
    if align is not None:
        p.alignment = align
    elif not justify:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p


def heading1(text):
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)


def heading2(text):
    p = doc.add_paragraph(style='Heading 2')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)


def page_break():
    if doc.paragraphs:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    else:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_figure(path, caption, width_in=5.6):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Inches(width_in))
        cp = doc.add_paragraph(style='Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.first_line_indent = Inches(0)
        cr = cp.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(10)
        cr.font.name = 'Times New Roman'
    else:
        add_para(f"[FIGURE NOT FOUND: {path}]")


def add_table(headers, rows, caption=None, col_widths=None, font_size=9.5, note=None):
    if caption:
        cp = doc.add_paragraph(style='Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.first_line_indent = Inches(0)
        cr = cp.add_run(caption)
        cr.bold = True
        cr.font.size = Pt(10.5)
        cr.font.name = 'Times New Roman'
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        t.style = 'Table Grid'
    except KeyError:
        pass
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    if col_widths:
        usable_in = (doc.sections[0].page_width - doc.sections[0].left_margin
                     - doc.sections[0].right_margin) / 914400
        cw = list(col_widths)
        if sum(cw) > usable_in:
            scale = usable_in / sum(cw)
            cw = [w * scale for w in cw]
        dxa = [int(round(w * 1440)) for w in cw]
        old_grid = tbl.find(qn('w:tblGrid'))
        if old_grid is not None:
            tbl.remove(old_grid)
        grid = OxmlElement('w:tblGrid')
        for w in dxa:
            gc = OxmlElement('w:gridCol')
            gc.set(qn('w:w'), str(w))
            grid.append(gc)
        tbl.insert(list(tbl).index(tblPr) + 1, grid)
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW')
                    tcPr.append(tcW)
                tcW.set(qn('w:w'), str(dxa[ci]))
                tcW.set(qn('w:type'), 'dxa')

    def write_cell(cell, text, bold=False, align='center'):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, v in (('top', 60), ('start', 80), ('bottom', 60), ('end', 80)):
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(v))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)
        v_align = OxmlElement('w:vAlign')
        v_align.set(qn('w:val'), 'center')
        tcPr.append(v_align)
        cell.text = ''
        para = cell.paragraphs[0]
        para.paragraph_format.first_line_indent = Inches(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        run = para.add_run(text)
        run.bold = bold
        run.font.size = Pt(font_size)
        run.font.name = 'Times New Roman'
        para.alignment = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]

    for i, h in enumerate(headers):
        write_cell(t.rows[0].cells[i], h, bold=True, align='center')
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            write_cell(t.rows[ri].cells[ci], str(val),
                       align='left' if ci == 0 else 'center')
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    if note:
        np_ = doc.add_paragraph()
        np_.paragraph_format.first_line_indent = Inches(0)
        np_.paragraph_format.space_before = Pt(2)
        np_.paragraph_format.space_after = Pt(8)
        np_.paragraph_format.line_spacing = 1.0
        nr = np_.add_run("Note: " + note)
        nr.italic = True
        nr.font.size = Pt(9)
        nr.font.name = 'Times New Roman'


# -------- OMML displayed equation --------

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def add_equation(omml_inner, label=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p_elem = p._p
    pPr = p_elem.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for val, pos in [('center', '4230'), ('right', '8460')]:
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), val)
        tab.set(qn('w:pos'), pos)
        tabs.append(tab)
    pPr.append(tabs)
    p_elem.append(parse_xml(f'<w:r xmlns:w="{W_NS}"><w:tab/></w:r>'))
    p_elem.append(parse_xml(f'<m:oMath xmlns:m="{M_NS}">{omml_inner}</m:oMath>'))
    if label:
        p_elem.append(parse_xml(
            f'<w:r xmlns:w="{W_NS}"><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:sz w:val="24"/></w:rPr><w:tab/><w:t xml:space="preserve">{label}</w:t></w:r>'))
    return p


def mr(t, italic=True):
    t = t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if not italic else ''
    return f'<m:r>{sty}<m:t xml:space="preserve">{t}</m:t></m:r>'


def msub(b, s):
    return f'<m:sSub><m:e>{b}</m:e><m:sub>{s}</m:sub></m:sSub>'


def msup(b, s):
    return f'<m:sSup><m:e>{b}</m:e><m:sup>{s}</m:sup></m:sSup>'


def msubsup(b, sb, sp):
    return f'<m:sSubSup><m:e>{b}</m:e><m:sub>{sb}</m:sub><m:sup>{sp}</m:sup></m:sSubSup>'


def mfenced(i):
    return f'<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{i}</m:e></m:d>'


def mnary_sum(sub, sup, body):
    return ('<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/>'
            '<m:subHide m:val="0"/><m:supHide m:val="0"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{body}</m:e></m:nary>')


# ============================================================
# TITLE BLOCK
# ============================================================

PAPER_TITLE = "Layered Oil-Price Pass-Through in India: Evidence from Wholesale, Retail Fuel, and Consumer Prices"

tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.first_line_indent = Inches(0)
tp.paragraph_format.space_before = Pt(0)
tp.paragraph_format.space_after = Pt(12)
tr = tp.add_run(PAPER_TITLE)
tr.bold = True
tr.font.name = 'Times New Roman'
tr.font.size = Pt(16)

au = doc.add_paragraph()
au.alignment = WD_ALIGN_PARAGRAPH.CENTER
au.paragraph_format.first_line_indent = Inches(0)
ar = au.add_run("Aniket Pandey")
ar.font.name = 'Times New Roman'
ar.font.size = Pt(11)

sup_ = doc.add_paragraph()
sup_.alignment = WD_ALIGN_PARAGRAPH.CENTER
sup_.paragraph_format.first_line_indent = Inches(0)
sr2 = sup_.add_run("Under the supervision of Prof. Shakti Kumar")
sr2.italic = True
sr2.font.name = 'Times New Roman'
sr2.font.size = Pt(11)

aff = doc.add_paragraph()
aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
aff.paragraph_format.first_line_indent = Inches(0)
aff.paragraph_format.space_after = Pt(12)
afr = aff.add_run("Jawaharlal Nehru University, New Delhi")
afr.font.name = 'Times New Roman'
afr.font.size = Pt(11)

# ABSTRACT
ab_h = doc.add_paragraph()
ab_h.paragraph_format.first_line_indent = Inches(0)
ab_h.alignment = WD_ALIGN_PARAGRAPH.LEFT
abh = ab_h.add_run("Abstract")
abh.bold = True
abh.font.size = Pt(11)
abh.font.name = 'Times New Roman'

abstract_text = (
    "This paper studies how global oil-price shocks pass through India's domestic price system. "
    "Instead of estimating only the effect of oil prices on headline CPI, it follows the shock "
    "across wholesale, retail fuel, fuel-sensitive consumer, and headline consumer price layers. "
    "Monthly log-difference ADL models are estimated with Newey-West HAC inference and bootstrap "
    "symmetry checks. The results show a layered attenuation pattern. Retail petrol responds "
    "strongly to Brent shocks. WPI Fuel and Power responds strongly to rupee oil shocks, but with "
    "a functional-form caveat. CPI Fuel and Light gives smaller bridge evidence, and headline WPI "
    "shows modest pass-through. At the headline CPI endpoint, positive pass-through is weak and "
    "statistically insignificant. Strong short-run asymmetry is not supported, except for a "
    "marginal retail petrol result. Overall, oil shocks remain visible in fuel layers but lose "
    "force in broad consumer inflation."
)
ap = add_para_math(abstract_text, indent=False, size=10.5)
ap.paragraph_format.line_spacing = 1.0
ap.paragraph_format.space_after = Pt(6)

kp = doc.add_paragraph()
kp.paragraph_format.first_line_indent = Inches(0)
kp.paragraph_format.line_spacing = 1.0
kr1 = kp.add_run("Keywords: ")
kr1.bold = True
kr1.font.name = 'Times New Roman'
kr1.font.size = Pt(10.5)
kr2 = kp.add_run("Oil-price pass-through; inflation; India; wholesale price index; consumer price index; retail fuel prices.")
kr2.font.name = 'Times New Roman'
kr2.font.size = Pt(10.5)

jp = doc.add_paragraph()
jp.paragraph_format.first_line_indent = Inches(0)
jp.paragraph_format.line_spacing = 1.0
jp.paragraph_format.space_after = Pt(10)
jr1 = jp.add_run("JEL codes: ")
jr1.bold = True
jr1.font.name = 'Times New Roman'
jr1.font.size = Pt(10.5)
jr2 = jp.add_run("C22; E31; Q43; F31.")
jr2.font.name = 'Times New Roman'
jr2.font.size = Pt(10.5)

# ============================================================
# 1. INTRODUCTION
# ============================================================

heading1("1. Introduction")

add_para_math(
    "India imports a large share of its crude oil and prices that crude in United States dollars. "
    "Domestic oil pressure therefore depends on the Brent price and on the rupee-dollar exchange "
    "rate, so a useful measure of the domestic oil shock is the rupee-denominated oil price, "
    "oil^{INR}_t = Brent^{USD}_t × INR_t/USD_t. When Brent rises or the rupee weakens, the rupee "
    "oil price rises, and the cost pressure should appear first in the layers closest to fuel and "
    "then, with attenuation, in broader price aggregates (Mandal et al., 2012; "
    "Bhanumurthy et al., 2012)."
)

add_para_math(
    "There is a familiar puzzle in this transmission. Petrol and diesel prices visibly track crude, "
    "and wholesale fuel inflation moves with rupee oil prices, but headline consumer inflation does "
    "not move one-for-one with oil. The Wholesale Price Index (WPI) is closer to producer cost "
    "pressure, while the Consumer Price Index (CPI) is the household-facing inflation index and is "
    "weighted heavily towards food, housing, services, health, and education. Fuel is only one "
    "component of the CPI basket, so even a strong oil shock can be diluted by the time it reaches "
    "headline CPI inflation."
)

add_para_math(
    "Most existing studies on Indian oil pass-through estimate a single equation linking oil to "
    "either WPI or CPI inflation (Mandal et al., 2012; Bhanumurthy et al., 2012; Pradeep, 2022). "
    "That approach gives one elasticity, but it does not show where the signal weakens. The "
    "contribution of this paper is to treat oil-price pass-through as a layered price-system "
    "problem rather than a single elasticity. The chain studied here moves from the rupee oil "
    "price to wholesale prices, from Brent to the Petroleum Planning and Analysis Cell (PPAC) "
    "Delhi retail petrol price, from PPAC retail petrol to CPI Fuel and Light, and from rupee oil "
    "to headline CPI. Each layer is estimated as a short-run asymmetric autoregressive "
    "distributed lag (ADL) model in monthly log differences."
)

add_para_math(
    "The institutional background also matters. Domestic fuel prices were not always fully linked "
    "to international prices, and reform episodes changed the way international oil movements "
    "entered petrol and diesel prices. Pradeep (2022) studies this reform setting and places "
    "asymmetry at the centre of the question. This paper uses that background more narrowly. It "
    "does not claim that deregulation caused the post-2010 estimates. It treats the reform period "
    "as useful context for reading stronger wholesale and retail fuel pass-through after pricing "
    "became more market-linked."
)

add_para_math(
    "The WPI and CPI distinction is useful for this question because the two indices measure "
    "different layers of the price system. WPI is closer to producer and wholesale cost pressure. "
    "CPI is the consumer-facing aggregate used for inflation targeting, but it contains many "
    "items whose monthly movements need not follow oil prices. A layered design therefore gives "
    "a more cautious interpretation than a single headline CPI equation. If oil is strong in fuel "
    "layers but weak in headline CPI, the correct conclusion is attenuation, not irrelevance."
)

add_para_math(
    "Inference uses Newey-West heteroskedasticity and autocorrelation consistent (HAC) standard "
    "errors (Newey & West, 1987), and the cumulative pass-through statistics CPT^{+} and CPT^{-} "
    "summarise the response to positive and negative shock terms. Bootstrap resampling provides a "
    "robustness check on asymmetry. Diagnostic tests cover serial correlation, functional form, "
    "and recursive stability, and a model-gate rule separates main, caveated, and excluded "
    "equations. "
    "Bai and Perron (2003) structural-break logic and the unit-root test battery, including KPSS "
    "checks (Kwiatkowski et al., 1992), are used to assess the differenced specification."
)

add_para_math(
    "The estimates suggest a layered attenuation pattern. Retail petrol shows the strongest direct "
    "fuel response, with CPT^{+} = 0.3459 (p < 0.001). WPI Fuel and Power shows CPT^{+} = 0.2866 "
    "(p < 0.001), but it is carried with the HAC-RESET caveat discussed below. CPI Fuel and Light, "
    "the consumer fuel bridge, shows CPT^{+} = 0.1777 (p = 0.0021). "
    "Headline WPI shows a statistically significant but modest CPT^{+} = 0.0301 (p = 0.0240). "
    "Headline CPI shows CPT^{+} = 0.0213 with p = 0.1220, which is positive in sign but not "
    "statistically significant at conventional levels. A formal Wald test rejects equality of the "
    "retail petrol and headline CPI positive pass-through estimates at F = 14.3499, p = 0.0002. "
    "Asymmetry is mostly weak, with only retail petrol showing a marginal CPT^{+} = CPT^{-} "
    "rejection at the ten per cent level (p = 0.0999)."
)

add_para("The remainder of the paper is organised as follows. Section 2 describes the data and "
         "variables. Section 3 sets out the methodology. Section 4 reports results across price "
         "layers. Section 5 covers robustness, diagnostics, and limitations. Section 6 discusses "
         "implications and concludes.")

# ============================================================
# 2. DATA AND VARIABLES
# ============================================================

heading1("2. Data and variables")

add_para_math(
    "The dataset is monthly and combines official Indian price series with international oil and "
    "exchange-rate data. Headline WPI and WPI Fuel and Power are from the Office of the Economic "
    "Adviser (OEA), chained to the 2011-12 base (Office of the Economic Adviser, 2017). The "
    "headline CPI source file is the OECD all-items India CPI series accessed through FRED "
    "(INDCPIALLMINMEI), while CPI Fuel and Light is the processed all-India MoSPI component "
    "series. Delhi retail petrol prices are from the Petroleum Planning and Analysis Cell "
    "(PPAC, 2024). Brent crude is from the World Bank Pink Sheet (World Bank, 2026). The CPI "
    "pipeline uses the FRED EXINUS file, and the WPI pipeline uses the updated FRED "
    "EXINUS_latest file. All series are monthly."
)

add_para_math(
    "The domestic oil shock is constructed as oil^{INR}_t = Brent^{USD}_t × INR_t/USD_t, where "
    "INR_t/USD_t is the monthly average exchange rate. This rupee oil price captures both the "
    "Brent component and the exchange-rate component of cost pressure on Indian importers and "
    "refiners. For the retail petrol model, the regressor is the Brent shock directly, because "
    "domestic retail fuel prices are administered with reference to international product prices "
    "and the rupee component is partially absorbed by tax and pricing rules. For the CPI Fuel and "
    "Light bridge model, the regressor is the PPAC petrol shock, because that is the price the "
    "fuel-sensitive consumer layer faces."
)

add_para_math(
    "The samples are deliberately allowed to differ by layer. The headline WPI series gives the "
    "longest window, from 1983-05 to 2026-03. WPI Fuel and Power begins later, in 1995-05. The "
    "PPAC Delhi petrol and headline CPI equations run from 2004-08 to 2024-12 and therefore give "
    "the main 20-year consumer-side window. CPI Fuel and Light begins only in 2011-05, so it is "
    "used as a bridge between retail fuel and consumer inflation rather than as a mandatory "
    "headline result. This is why the title does not state one common date range for every layer."
)

add_equation(
    f"{msub(mr('Δx'), mr('t'))} {mr('=', italic=False)} {mr('100', italic=False)} "
    f"{mr('×', italic=False)} {mfenced(mr('ln') + mfenced(msub(mr('x'), mr('t'))) + mr('−', italic=False) + mr('ln') + mfenced(msub(mr('x'), mr('t−1'))))}",
    label="(1)"
)

add_para_math(
    "Equation (1) defines the monthly log difference of any series x_t. The factor of 100 means "
    "coefficients on Δx_t can be read approximately as percentage responses. The positive and "
    "negative shock components are defined as Δx^{+}_t = max(Δx_t, 0) and Δx^{-}_t = "
    "min(Δx_t, 0). The negative-shock variable is non-positive by construction, so the "
    "cumulative pass-through CPT^{-}, which is the sum of coefficients on lagged Δx^{-} terms, "
    "should not be interpreted as the response to a positive shock with the opposite sign."
)

add_table(
    headers=["Series", "Source", "Transformation", "Sample", "Role"],
    rows=[
        ["Brent crude price", "World Bank Pink Sheet monthly", "Monthly log difference", "Matched to layer", "Global oil shock"],
        ["INR/USD exchange rate", "FRED EXINUS / EXINUS_latest", "Monthly log difference", "Matched to layer", "Exchange-rate component"],
        ["Rupee oil price", "Brent × FRED INR/USD", "Monthly log difference", "Matched to layer", "Domestic oil shock"],
        ["Headline WPI", "OEA WPI files, base 2011-12", "Monthly log difference", "1983-05 to 2026-03", "Wholesale endpoint"],
        ["WPI Fuel and Power", "OEA WPI files, base 2011-12", "Monthly log difference", "1995-05 to 2026-03", "Wholesale fuel layer"],
        ["PPAC Delhi retail petrol", "PPAC RSP pre/post-2017 files", "Monthly log difference", "2004-08 to 2024-12", "Retail fuel layer"],
        ["CPI Fuel and Light", "MoSPI all-India Fuel and Light", "Monthly log difference", "2011-05 to 2024-12", "Consumer fuel bridge"],
        ["Headline CPI", "FRED/OECD INDCPIALLMINMEI", "Monthly log difference", "2004-08 to 2024-12", "Consumer endpoint"],
    ],
    caption="Table 1. Data, transformations, and model role.",
    col_widths=[1.4, 1.5, 1.4, 1.4, 1.0],
    note="All transformations are multiplied by 100, so coefficients can be read approximately as percentage responses. Samples differ because the official series are available over different periods. Source: Author's compilation from cited official sources."
)

# ============================================================
# 3. METHODOLOGY
# ============================================================

heading1("3. Methodology")

add_para_math(
    "Each price layer is modelled with a short-run asymmetric ADL specification in monthly log "
    "differences. Let Δy_t denote the monthly log change of the layer-specific dependent variable "
    "and Δx_t denote the relevant shock variable. The estimating equation is"
)

eq_body = (
    f"{msub(mr('Δy'), mr('t'))} {mr('=', italic=False)} {mr('α')} "
    f"{mr('+', italic=False)} "
    + mnary_sum(msub(mr('i'), mr('=1')), mr('p'),
                msub(mr('φ'), mr('i')) + msub(mr('Δy'), mr('t−i')))
    + f" {mr('+', italic=False)} "
    + mnary_sum(msub(mr('j'), mr('=0')), mr('q'),
                msubsup(mr('β'), mr('j'), mr('+')) + msubsup(mr('Δx'), mr('t−j'), mr('+')))
    + f" {mr('+', italic=False)} "
    + mnary_sum(msub(mr('j'), mr('=0')), mr('q'),
                msubsup(mr('β'), mr('j'), mr('−')) + msubsup(mr('Δx'), mr('t−j'), mr('−')))
    + f" {mr('+', italic=False)} {msup(mr('γ'), mr('′'))}{msub(mr('Z'), mr('t'))} "
    + f"{mr('+', italic=False)} {msub(mr('μ'), mr('m'))} "
    + f"{mr('+', italic=False)} {msub(mr('ε'), mr('t'))}"
)
add_equation(eq_body, label="(2)")

add_para_math(
    "Here Δy_t is monthly inflation in the relevant price index and Δx_t is the layer-appropriate "
    "shock. The lag choices are model-specific and are fixed before interpreting the results. In "
    "the CPI pipeline, the autoregressive lag is p = 3, selected by AIC over p = 1 to 4, and the "
    "oil-shock lag window is q = 3. In the WPI pipeline, the headline and Fuel and Power models "
    "use AR(12) terms and oil-shock lags 0 to 6. The vector Z_t contains only the controls used "
    "in each estimated equation, and μ_m denotes month fixed effects where they are included. "
    "The innovation term ε_t is allowed to be heteroskedastic and serially correlated."
)

add_para_math(
    "The exact specifications used for the main and caveated rows are: headline WPI, ADL(12,6) "
    "with month fixed effects; "
    "WPI Fuel and Power, ADL(12,6) with exchange-rate controls, reform and COVID dummies, and "
    "month fixed effects; PPAC retail petrol, ADL(3,3) with IIP growth, petrol and diesel "
    "deregulation dummies, a COVID dummy, and month fixed effects; PPAC petrol to CPI Fuel and "
    "Light, ADL(3,3) without extra controls; and headline CPI M1, ADL(3,3) with IIP growth, "
    "petrol and diesel deregulation dummies, a COVID dummy, and month fixed effects. Full "
    "coefficient outputs are retained as supplementary model tables: WPI tables 04b and 06b, "
    "CPI table 06, PPAC table 22b, and bridge table 27b. The article reports cumulative "
    "pass-through to stay within the page limit."
)

add_para_math(
    "Variables across layers are matched to the channel each layer represents. The headline WPI "
    "model uses rupee oil shocks as Δx_t. The WPI Fuel and Power model also uses rupee oil shocks. "
    "The PPAC retail petrol model uses Brent shocks. The CPI Fuel and Light bridge model uses "
    "PPAC petrol shocks. The headline CPI model uses rupee oil shocks. This matching keeps each "
    "regression interpretable as a pass-through equation for its layer."
)

add_para_math(
    "Cumulative pass-through is summarised by CPT^{+} = Σ_j β^{+}_j and CPT^{-} = Σ_j β^{-}_j. "
    "Hypothesis tests are reported for CPT^{+} = 0, CPT^{-} = 0, and CPT^{+} = CPT^{-}. Inference "
    "uses Newey-West HAC standard errors with bandwidth floor(0.75N^{1/3}) (Newey & West, 1987). "
    "A restricted-residual circular block bootstrap with 4,999 replications provides a "
    "robustness check on the asymmetry test. Predictive precedence between oil and the dependent "
    "variable is examined with Granger-style tests on lagged shock terms; these are reported as "
    "predictive precedence rather than structural causality. Structural-break behaviour is "
    "examined through Bai and Perron (2003) tests, and stationarity of log-differenced series is "
    "verified by KPSS tests (Kwiatkowski et al., 1992)."
)

add_para_math(
    "A model is treated as claim-bearing only if it passes a mandatory diagnostic gate, "
    "specifically a Breusch-Godfrey test for residual serial correlation, a HAC-robust RESET test "
    "for functional form, and a recursive cumulative sum test for parameter stability. Models "
    "that fail one of these checks are reported with an explicit caveat or excluded from the main "
    "claims. ADF and Phillips-Perron tests support the first-differenced design; KPSS passes the "
    "CPI differenced variables and most WPI differenced variables, while WPI inflation carries a "
    "KPSS break-related caveat handled through the Bai-Perron check. The framework is short-run "
    "by design. There are no NARDL bounds tests, no "
    "error-correction terms, and no long-run pass-through claims, because the object of interest "
    "is monthly transmission rather than long-run equilibrium adjustment."
)

# ============================================================
# 4. RESULTS
# ============================================================

heading1("4. Results")

heading2("4.1 Cumulative pass-through across price layers")

add_para_math(
    "Table 2 reports the central estimates. The reading is layer by layer. The retail petrol "
    "model uses Brent shocks over 2004-08 to 2024-12. The wholesale models use rupee oil shocks "
    "over the longest available samples. The CPI Fuel and Light bridge uses PPAC petrol shocks "
    "over 2011-05 to 2024-12. The headline CPI model uses rupee oil shocks over 2004-08 to "
    "2024-12. Sample sizes vary across layers because the official series are available over "
    "different periods, so the full five-layer ranking is descriptive rather than a single "
    "common-sample estimate."
)

add_table(
    headers=["Layer", "Sample", "N", "CPT+", "p", "CPT-", "p", "Asym. p", "Verdict"],
    rows=[
        ["PPAC retail petrol", "2004-08 to 2024-12", "245", "0.3459", "<0.001", "0.1912", "0.0002", "0.0999", "Strong direct fuel; marginal asymmetry"],
        ["WPI Fuel and Power", "1995-05 to 2026-03", "371", "0.2866", "<0.001", "0.2677", "<0.001", "0.7832", "Strong wholesale fuel; RESET caveat"],
        ["CPI Fuel and Light", "2011-05 to 2024-12", "164", "0.1777", "0.0021", "0.1058", "0.1741", "0.4554", "Bridge evidence; shorter sample"],
        ["Headline WPI", "1983-05 to 2026-03", "515", "0.0301", "0.0240", "0.0374", "0.0012", "0.6727", "Modest but significant"],
        ["Headline CPI", "2004-08 to 2024-12", "245", "0.0213", "0.1220", "0.0006", "0.9375", "0.2408", "Weak; not significant"],
    ],
    caption="Table 2. Cumulative pass-through by price layer.",
    col_widths=[1.35, 1.2, 0.4, 0.55, 0.55, 0.55, 0.55, 0.55, 1.6],
    font_size=9,
    note="CPT+ and CPT- are sums of positive- and negative-shock coefficients across the ADL lag window. P-values use Newey-West HAC inference. Asym. p is the p-value for the test CPT+ = CPT-. WPI models use rupee oil shocks; the retail petrol model uses Brent shocks; the CPI Fuel and Light row is the PPAC-to-Fuel bridge. Samples differ by data availability, so the table is descriptive; the consumer-chain common-sample check is reported below. Source: Author's estimates from model outputs."
)

heading2("4.2 Wholesale layer")

add_para_math(
    "Headline WPI shows statistically significant but modest pass-through from rupee oil shocks. "
    "The cumulative response to positive shocks is CPT^{+} = 0.0301 with p = 0.0240, and the "
    "cumulative response to negative shock terms is CPT^{-} = 0.0374 with p = 0.0012. The "
    "asymmetry test does not reject equality (p = 0.6727), so the wholesale endpoint does not "
    "support short-run asymmetry. The adjusted R^{2} of 0.4207 over 515 observations indicates "
    "that the lagged inflation, oil shock, and seasonal structure together explain a substantial "
    "share of monthly headline WPI variation. The model passes the Breusch-Godfrey, HAC-RESET, "
    "and recursive CUSUM checks, so it is treated as claim-bearing."
)

add_para_math(
    "WPI Fuel and Power carries a much stronger oil signal, with CPT^{+} = 0.2866 (p < 0.001) and "
    "CPT^{-} = 0.2677 (p < 0.001) over 371 observations. The point estimates imply that "
    "approximately 28 to 29 per cent of a one per cent rupee oil shock passes into the wholesale "
    "fuel layer in the months captured by the lag window. The asymmetry p-value of 0.7832 is "
    "high, so positive and negative effects are statistically indistinguishable. The "
    "Breusch-Godfrey and recursive CUSUM tests pass, but the HAC-RESET test for functional form "
    "fails. The result is therefore reported as mechanism evidence with a functional-form caveat, "
    "rather than as a fully validated claim-bearing endpoint."
)

add_para_math(
    "Figure 1 compares the wholesale pass-through estimates visually. The difference between "
    "headline WPI and WPI Fuel and Power is large. This is not surprising, but it is important "
    "for interpretation. It means that a broad wholesale aggregate can show only a small average "
    "response even when the fuel component inside the same price system is strongly exposed to "
    "oil movements. The figure also helps separate economic size from statistical significance. "
    "Headline WPI is statistically significant, but the coefficient is small. WPI Fuel and Power "
    "is both statistically significant and economically larger, although it carries the RESET "
    "caveat noted above."
)

add_figure(
    os.path.join(ROOT, "models/wpi/outputs/figures/fig_04_cumulative_passthrough.png"),
    "Figure 1. Cumulative oil pass-through in headline WPI and WPI Fuel and Power. Source: Author's estimates from WPI model outputs.",
    width_in=5.45
)

heading2("4.3 Retail petrol and the consumer fuel bridge")

add_para_math(
    "The PPAC Delhi retail petrol equation is the strongest direct fuel layer in the system. "
    "With Brent as the shock variable, CPT^{+} = 0.3459 (p < 0.001) and CPT^{-} = 0.1912 "
    "(p = 0.0002) over 245 observations. The asymmetry p-value is 0.0999, which rejects equality "
    "only at the ten per cent level. The result should therefore be read as suggestive of "
    "stronger upward than downward pass-through in retail petrol, but not as decisive evidence of "
    "asymmetric pricing. The retail petrol model passes the mandatory mechanism gate, so it can "
    "carry the channel claim."
)

add_para_math(
    "CPI Fuel and Light is the consumer fuel layer that sits between retail fuel prices and "
    "broad consumer inflation. With the PPAC petrol shock as the regressor, CPT^{+} = 0.1777 "
    "(p = 0.0021) and CPT^{-} = 0.1058 (p = 0.1741) over 164 observations. The positive "
    "pass-through is statistically significant; the negative side is not. The result is treated "
    "as bridge evidence rather than as a primary consumer-price result, because the sample only "
    "begins in 2011-05 and is therefore considerably shorter than the headline CPI sample. The "
    "estimate is consistent with retail fuel pressure entering the fuel-sensitive consumer layer "
    "with attenuation."
)

add_para_math(
    "The contrast between these two consumer-side layers is useful. Retail petrol is a narrow, "
    "policy-sensitive retail price and can move sharply when international prices are passed on. CPI "
    "Fuel and Light is already a basket component. It includes household fuel items and therefore "
    "smooths the retail petrol shock. The decline from 0.3459 in the retail petrol equation to "
    "0.1777 in the bridge equation is consistent with partial transmission rather than a full "
    "one-for-one movement. Because the bridge sample is short, the estimate is not used to make a "
    "20-year headline claim."
)

heading2("4.4 Headline CPI endpoint and integrated attenuation")

add_para_math(
    "Headline CPI is where the oil signal becomes weak. The positive cumulative pass-through is "
    "CPT^{+} = 0.0213 with p = 0.1220, and the negative side is CPT^{-} = 0.0006 with p = 0.9375. "
    "The positive coefficient has the expected sign, but it is not statistically significant at "
    "conventional levels. Asymmetry is not rejected (p = 0.2408). The model passes the mandatory "
    "M1 diagnostic gate, so the failure to find significant pass-through is an empirical result, "
    "not an artefact of a rejected specification."
)

add_para_math(
    "The ranking of positive pass-through estimates gives the central result. Retail petrol has "
    "the largest CPT^{+}, followed by WPI Fuel and Power and CPI Fuel and Light. The broad "
    "headline indices are much smaller, especially headline CPI. The formal Wald test for "
    "equality of retail petrol and headline CPI positive pass-through rejects the null at "
    "F = 14.3499, p = 0.0002. The evidence is therefore best read as attenuation across price "
    "layers rather than absence of oil-price transmission. Because not every layer shares the "
    "same sample window, the ranking is supported most directly for the consumer chain. Figure 2 "
    "visualises that chain."
)

add_figure(
    os.path.join(ROOT, "models/cpi/outputs/figures/fig_13_dilution_chain.png"),
    "Figure 2. Cumulative positive pass-through across the consumer-price chain. Source: Author's estimates from PPAC and CPI model outputs."
)

add_para_math(
    "A common-sample exercise re-estimates the consumer chain on the overlapping consumer-side "
    "window, with small N differences caused by lag construction, to reduce the risk that sample "
    "composition drives the ranking. Stage 1 (Brent to PPAC petrol) gives "
    "CPT^{+} = 0.4007 (p = 0.0001). Stage 2 (PPAC petrol to CPI Fuel and Light) gives "
    "CPT^{+} = 0.1777 (p = 0.0021). Stage 3 (rupee oil to headline CPI) gives CPT^{+} = 0.0064 "
    "(p = 0.6355). The qualitative pattern survives the common-sample restriction."
)

add_para_math(
    "Predictive-precedence tests give supporting evidence, but they are not used as structural "
    "causality tests. In the WPI system, lagged rupee oil changes predict headline WPI inflation "
    "(F = 7.4288, p < 0.001) and WPI Fuel and Power inflation (F = 15.2231, p < 0.001), while the "
    "reverse WPI-to-oil direction is not supported. On the consumer side, Brent strongly predicts "
    "PPAC petrol (F = 10.7127, p < 0.001). Lagged rupee oil terms predict headline CPI only at "
    "the ten per cent level (F = 2.2861, p = 0.0794), and PPAC petrol does not clearly predict CPI "
    "Fuel and Light in the bridge equation (p = 0.1126). This pattern matches the main estimates: "
    "the oil signal is easier to detect near fuel prices than in broad CPI."
)

# ============================================================
# 5. ROBUSTNESS, DIAGNOSTICS, AND LIMITATIONS
# ============================================================

heading1("5. Robustness, diagnostics, and limitations")

add_para_math(
    "The headline WPI result is robust to decomposing the rupee oil shock into a Brent component "
    "and an exchange-rate component. The decomposition gives CPT^{+} = 0.0309 (p = 0.0234) and "
    "CPT^{-} = 0.0375 (p < 0.001), with an asymmetry p-value of 0.7039. These figures are within "
    "rounding distance of the rupee-shock benchmark and confirm that the modest wholesale "
    "endpoint result does not depend on combining the two components. Bootstrap symmetry tests "
    "with 4,999 replications fail to reject CPT^{+} = CPT^{-} for headline WPI and for WPI Fuel "
    "and Power, which is consistent with the HAC-based asymmetry tests in Table 2."
)

add_para_math(
    "The diagnostics are summarised in Table 3. The headline WPI model passes the Breusch-Godfrey "
    "test for serial correlation, the HAC-RESET test for functional form, and the recursive CUSUM "
    "test for parameter stability. The WPI Fuel and Power model passes Breusch-Godfrey and "
    "recursive CUSUM but fails HAC-RESET. The headline CPI M1 model passes the mandatory gate. "
    "The PPAC retail petrol model passes the mechanism gate. Two alternative CPI specifications "
    "explored in the supporting work, M2 and M3, are not claim-bearing because their diagnostics "
    "are rejected. The bootstrap symmetry checks do not reject symmetry in the main wholesale and "
    "consumer endpoint models, so asymmetry is not the central finding of the paper. Unit-root "
    "tests support using log differences, with the WPI inflation KPSS caveat read alongside the "
    "Bai-Perron break evidence."
)

add_table(
    headers=["Check", "Result", "Interpretation"],
    rows=[
        ["Headline WPI diagnostics", "BG, HAC-RESET, Rec-CUSUM pass", "Main WPI model accepted"],
        ["WPI Fuel and Power diagnostics", "BG, Rec-CUSUM pass; HAC-RESET fails", "Mechanism with caveat"],
        ["Headline CPI M1 gate", "Pass", "Main CPI endpoint accepted"],
        ["PPAC mechanism gate", "Pass", "Direct fuel channel accepted"],
        ["CPI M2, M3 gates", "Diagnostics fail", "Excluded from main claims"],
        ["Bootstrap symmetry (WPI, CPI)", "Symmetry not rejected", "Asymmetry not central"],
    ],
    caption="Table 3. Diagnostics and robustness summary.",
    col_widths=[2.1, 2.4, 2.0],
    font_size=10,
    note="BG = Breusch-Godfrey; HAC-RESET = HAC-robust Ramsey RESET; Rec-CUSUM = recursive cumulative sum stability test. Source: Author's diagnostic outputs."
)

add_para_math(
    "A pre/post-2010 split of the wholesale models offers a fourth piece of context. For headline "
    "WPI, CPT^{+} rises from 0.0117 (p = 0.3812) in the pre-2010 sample to 0.0741 (p = 0.0043) in "
    "the post-2010 sample. For WPI Fuel and Power, CPT^{+} rises from 0.0922 (p = 0.1679) to "
    "0.5241 (p < 0.001). The post-2010 estimates are consistent with stronger pass-through "
    "around the period of more market-linked fuel pricing, including petrol deregulation in "
    "June 2010 and diesel deregulation in October 2014. The split is not a clean causal estimate "
    "of deregulation, because other macroeconomic and policy changes occurred over the same "
    "window. It is reported as institutional evidence."
)

add_para_math(
    "The limitations are stated plainly. The estimates are reduced-form short-run projections, "
    "not structural causal effects. Different layers have different sample spans, although the "
    "common-sample exercise reproduces the qualitative ranking. CPI Fuel and Light has a shorter "
    "sample than headline CPI and WPI, and is therefore used as bridge rather than primary "
    "evidence. PPAC retail petrol uses Delhi prices, which are treated as a retail fuel-channel "
    "proxy rather than a complete national retail measure. The pre/post-2010 split is suggestive "
    "institutional evidence, not a clean deregulation experiment. The layered map is not a single "
    "mechanical WPI-to-CPI causal chain; it is a set of layer-by-layer pass-through equations "
    "that share a common interpretation."
)

# ============================================================
# 6. DISCUSSION AND CONCLUSION
# ============================================================

heading1("6. Discussion and conclusion")

add_para_math(
    "The result of this paper is not that oil is unimportant for Indian inflation. Retail petrol "
    "tracks Brent strongly, the wholesale fuel layer responds firmly to rupee oil shocks, and CPI "
    "Fuel and Light shows a smaller but statistically significant bridge response. The result is "
    "that headline CPI absorbs and dilutes oil shocks because the CPI basket is broad. Food, "
    "housing, services, health, and education together dominate the index, and the direct fuel "
    "weight, even augmented by indirect cost effects, is too small for the oil signal to dominate "
    "the headline figure on a monthly basis."
)

add_para_math(
    "For inflation monitoring, the implication is narrow but useful. A reader who monitors only "
    "headline CPI may miss near-fuel price pressure, because the oil signal is small at that "
    "layer and is statistically insignificant in the present sample. Monitoring a small set of "
    "fuel-sensitive layers, such as retail petrol, WPI Fuel and Power, and CPI Fuel and Light, "
    "carries more information about short-run pass-through than the headline aggregate alone. "
    "This is consistent with earlier work on Indian oil pass-through (Mandal et al., 2012; "
    "Bhanumurthy et al., 2012; Pradeep, 2022), and adds the layered organisation as a way to read "
    "those results."
)

add_para_math(
    "The conclusion can be stated compactly. The research question is whether oil-price shocks "
    "transmit uniformly across Indian price layers or whether they attenuate as the index becomes "
    "broader. The estimates support a descriptive attenuation ranking: retail petrol > WPI Fuel "
    "and Power > CPI Fuel and Light > headline WPI > headline CPI in CPT^{+} terms, with the "
    "formal consumer-chain test rejecting equality between retail petrol and headline CPI. The "
    "WPI Fuel and Power estimate remains mechanism evidence with a functional-form caveat. "
    "Asymmetry is not the main finding; only retail petrol shows a marginal asymmetry result at "
    "the ten per cent level. The post-2010 split is consistent with stronger pass-through under "
    "more market-linked fuel pricing, but it is institutional rather than causal evidence."
)

add_para("The main lesson is that oil shocks do not disappear in India, but they lose force as "
         "they move from fuel prices to broad consumer inflation.")

# ============================================================
# REFERENCES (APA)
# ============================================================

heading1("References")

REFS = [
    "Bai, J., & Perron, P. (2003). Computation and analysis of multiple structural change models. Journal of Applied Econometrics, 18(1), 1\u201122. https://doi.org/10.1002/jae.659",
    "Bhanumurthy, N. R., Das, S., & Bose, S. (2012). Oil price shock, pass-through policy and its impact on India (NIPFP Working Paper No. 2012-99). National Institute of Public Finance and Policy.",
    "Kwiatkowski, D., Phillips, P. C. B., Schmidt, P., & Shin, Y. (1992). Testing the null hypothesis of stationarity against the alternative of a unit root. Journal of Econometrics, 54(1\u20113), 159\u2011178. https://doi.org/10.1016/0304-4076(92)90104-Y",
    "Mandal, K., Bhattacharyya, I., & Bhoi, B. B. (2012). Is the oil price pass-through in India any different? Journal of Policy Modeling, 34(6), 832\u2011848. https://doi.org/10.1016/j.jpolmod.2012.06.001",
    "Ministry of Statistics and Programme Implementation. (2015). Consumer Price Index: Changes in the revised series. Government of India.",
    "Newey, W. K., & West, K. D. (1987). A simple, positive semi-definite, heteroskedasticity and autocorrelation consistent covariance matrix. Econometrica, 55(3), 703\u2011708. https://doi.org/10.2307/1913610",
    "Office of the Economic Adviser. (2017). Manual on Wholesale Price Index: Base 2011\u201112 = 100. Department for Promotion of Industry and Internal Trade, Ministry of Commerce and Industry, Government of India.",
    "Petroleum Planning and Analysis Cell. (2024). Ready reckoner: India's oil and gas. Ministry of Petroleum and Natural Gas, Government of India.",
    "Pradeep, S. (2022). Impact of diesel price reforms on asymmetricity of oil price pass-through to inflation: Indian perspective. The Journal of Economic Asymmetries, 26, e00249. https://doi.org/10.1016/j.jeca.2022.e00249",
    "World Bank. (2026). Commodity price data: The Pink Sheet. World Bank Commodity Markets.",
]

for ref in REFS:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(ref)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)


doc.save(OUT)
print(f"Saved: {OUT}")
