"""Build aniket-dissertation.docx strictly following blueprint.md.

ADL-only, layered attenuation story, plain MS Economics student prose, no em
dashes, APA 7 references, tables and figures inside chapters. Target ~7,000
words.
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
import os

ROOT = "/Users/aniketpandey/Desktop/fresh-dissertation"
OUT = os.path.join(ROOT, "aniket-dissertation.docx")

doc = Document()

for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Cm(3.54)
    section.right_margin = Inches(1)
    sectPr = section._sectPr
    for existing in sectPr.findall(qn('w:mirrorMargins')):
        sectPr.remove(existing)
    pgMar = sectPr.find(qn('w:pgMar'))
    if pgMar is not None:
        pgMar.set(qn('w:gutter'), '0')

settings_el = doc.settings.element
for existing in settings_el.findall(qn('w:mirrorMargins')):
    settings_el.remove(existing)


def force_tnr(style):
    style.font.name = 'Times New Roman'
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(k), 'Times New Roman')


normal = doc.styles['Normal']
normal.font.size = Pt(12)
force_tnr(normal)
pf = normal.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(6)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, sz, sp_b, sp_a in [('Heading 1', 20, 24, 18), ('Heading 2', 14, 14, 6), ('Heading 3', 12, 10, 4)]:
    s = doc.styles[name]
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    force_tnr(s)
    s.paragraph_format.space_before = Pt(sp_b)
    s.paragraph_format.space_after = Pt(sp_a)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing = 1.5
doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

cs = doc.styles['Caption']
cs.font.size = Pt(10)
cs.font.color.rgb = RGBColor(0, 0, 0)
force_tnr(cs)
cs.paragraph_format.line_spacing = 1.0


def add_para(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p


def add_para_math(text):
    """Add a paragraph with lightweight inline sub/superscript parsing.

    This mirrors the older builder: _x, _{xx}, ^x, and ^{xx} become Word
    subscript/superscript runs. Use braces for multi-character subscripts.
    """
    p = doc.add_paragraph()
    i = 0
    n = len(text)
    while i < n:
        marker = text[i]
        if marker in ('_', '^'):
            subscript = marker == '_'
            i += 1
            if i < n and text[i] == '{':
                j = text.find('}', i)
                if j == -1:
                    content = text[i:]
                    i = n
                else:
                    content = text[i + 1:j]
                    i = j + 1
            elif i < n:
                content = text[i]
                i += 1
            else:
                content = ''
            r = p.add_run(content)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.subscript = subscript
            r.font.superscript = not subscript
        else:
            j = i
            while j < n and text[j] not in ('_', '^'):
                j += 1
            r = p.add_run(text[i:j])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            i = j
    return p


def chapter_heading(num, title):
    p = doc.add_paragraph(style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if num:
        r1 = p.add_run(num)
        r1.bold = True
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(20)
        r1.add_break()
    r2 = p.add_run(title)
    r2.bold = True
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(20)


def section_heading(text):
    p = doc.add_paragraph(style='Heading 2')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)


def page_break():
    if doc.paragraphs:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    else:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_figure(path, caption, width_in=5.5):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        pic = p.add_run().add_picture(path, width=Inches(width_in))
        pic._inline.docPr.set('title', os.path.basename(path))
        pic._inline.docPr.set('descr', caption)
        cp = doc.add_paragraph(style='Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(10)
        cr.font.name = 'Times New Roman'
    else:
        add_para(f"[FIGURE NOT FOUND: {path}]")


def add_table(headers, rows, caption=None, col_widths=None, font_size=10):
    if caption:
        cp = doc.add_paragraph(style='Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cr = cp.add_run(caption)
        cr.bold = True
        cr.font.size = Pt(11)
        cr.font.name = 'Times New Roman'
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        t.style = 'Table Grid'
    except KeyError:
        pass
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    if col_widths:
        usable_width_in = (
            doc.sections[0].page_width
            - doc.sections[0].left_margin
            - doc.sections[0].right_margin
        ) / 914400
        col_widths = list(col_widths)
        if sum(col_widths) > usable_width_in:
            scale = usable_width_in / sum(col_widths)
            col_widths = [w * scale for w in col_widths]
        dxa_widths = [int(round(w * 1440)) for w in col_widths]
        total = sum(dxa_widths)
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        tblW.set(qn('w:w'), str(total))
        tblW.set(qn('w:type'), 'dxa')
        old_grid = tbl.find(qn('w:tblGrid'))
        if old_grid is not None:
            tbl.remove(old_grid)
        grid = OxmlElement('w:tblGrid')
        for w in dxa_widths:
            gc = OxmlElement('w:gridCol')
            gc.set(qn('w:w'), str(w))
            grid.append(gc)
        tbl.insert(list(tbl).index(tblPr) + 1, grid)
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW')
                    tcPr.append(tcW)
                tcW.set(qn('w:w'), str(dxa_widths[ci]))
                tcW.set(qn('w:type'), 'dxa')
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        para = hdr[i].paragraphs[0]
        run = para.add_run(h)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.name = 'Times New Roman'
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
    for ri, row in enumerate(rows, start=1):
        cells = t.rows[ri].cells
        for ci, val in enumerate(row):
            cells[ci].text = ''
            para = cells[ci].paragraphs[0]
            run = para.add_run(str(val))
            run.font.size = Pt(font_size)
            run.font.name = 'Times New Roman'
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(0)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    return t


def add_note(text):
    p = doc.add_paragraph()
    r = p.add_run("Notes: " + text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.name = 'Times New Roman'
    p.paragraph_format.space_after = Pt(10)


M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def add_equation(omml_inner, label=None):
    p = doc.add_paragraph()
    p_elem = p._p
    pPr = p_elem.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for val, pos in [('center', '4230'), ('right', '8460')]:
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), val)
        tab.set(qn('w:pos'), pos)
        tabs.append(tab)
    pPr.append(tabs)
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), '120')
    sp.set(qn('w:after'), '120')
    sp.set(qn('w:line'), '360')
    sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)
    p_elem.append(parse_xml(f'<w:r xmlns:w="{W_NS}"><w:tab/></w:r>'))
    p_elem.append(parse_xml(f'<m:oMath xmlns:m="{M_NS}">{omml_inner}</m:oMath>'))
    if label:
        p_elem.append(parse_xml(
            f'<w:r xmlns:w="{W_NS}"><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:sz w:val="24"/></w:rPr><w:tab/><w:t xml:space="preserve">{label}</w:t></w:r>'))
    return p


def mr(t):
    t = t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    return f'<m:r><m:t xml:space="preserve">{t}</m:t></m:r>'


def msub(b, s):
    return f'<m:sSub><m:e>{b}</m:e><m:sub>{s}</m:sub></m:sSub>'


def msup(b, s):
    return f'<m:sSup><m:e>{b}</m:e><m:sup>{s}</m:sup></m:sSup>'


def mfenced(i):
    return f'<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{i}</m:e></m:d>'


def mnary_sum(sub, sup, body):
    return ('<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/>'
            '<m:subHide m:val="0"/><m:supHide m:val="0"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{body}</m:e></m:nary>')


# Title page
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(120)
tr = tp.add_run("From wholesale prices to consumer inflation: layered pass-through of global oil shocks in India, 1983-2026")
tr.bold = True
tr.font.name = 'Times New Roman'
tr.font.size = Pt(20)

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.paragraph_format.space_before = Pt(24)
sr = sp.add_run("Evidence from short-run asymmetric ADL models")
sr.italic = True
sr.font.name = 'Times New Roman'
sr.font.size = Pt(14)
page_break()

# Chapters loaded from supplementary modules
import _diss_chapters
_diss_chapters.build(doc, add_para, chapter_heading, section_heading, page_break,
                     add_figure, add_table, add_note, add_equation,
                     add_para_math,
                     mr, msub, msup, mfenced, mnary_sum, ROOT)

doc.save(OUT)
print(f"Saved: {OUT}")
