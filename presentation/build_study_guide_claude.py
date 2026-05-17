"""Build dissertation_study_guide-claude.docx (and PDF).

A beginner-friendly, explanatory walk-through of Aniket Pandey's dissertation
and journal paper on layered oil-price pass-through in India. Written so a
reader who has not touched econometrics in a while can still follow every
step, every variable, every test, and every result, and answer viva-style
questions.

The guide is split into part files so each section stays manageable:
    _guide_part1_overview.py        Big picture, why, what, one-page map
    _guide_part2_background.py      India, WPI, CPI, dollar oil, deregulation
    _guide_part3_research_design.py Research question, objectives, contribution
    _guide_part4_data.py            Data sources with full forms and roles
    _guide_part5_variables.py       Log-difference, shock split, rupee oil
    _guide_part6_methodology.py     ADL, OLS, Greek letters, ADL(p,q)
    _guide_part7_inference.py       HAC, bootstrap, Wald, Granger in plain terms
    _guide_part8_diagnostics.py     Unit root, break, BG, RESET, CUSUM, gate
    _guide_part9_results.py         Layer-by-layer results, numbers explained
    _guide_part10_robustness.py     Robustness, subsamples, limitations
    _guide_part11_conclusion.py     Conclusion, policy read, takeaways
    _guide_part12_qa.py             Likely viva questions with model answers
    _guide_part13_glossary.py       Glossary and cheat-sheet

Output:
    /Users/aniketpandey/Desktop/fresh-dissertation/dissertation_study_guide-claude.docx
    /Users/aniketpandey/Desktop/fresh-dissertation/dissertation_study_guide-claude.pdf (via soffice)
"""

import os
import sys
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

OUT_DOCX = ROOT / "dissertation_study_guide-claude.docx"
OUT_PDF = ROOT / "dissertation_study_guide-claude.pdf"

doc = Document()

# --- Page setup -------------------------------------------------------------
for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    run = paragraph.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


for section in doc.sections:
    _add_page_number(section.footer.paragraphs[0])


# --- Force Times New Roman everywhere --------------------------------------
def force_tnr(style):
    style.font.name = 'Times New Roman'
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(k), 'Times New Roman')


normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
force_tnr(normal)
pf = normal.paragraph_format
pf.line_spacing = 1.3
pf.space_before = Pt(2)
pf.space_after = Pt(4)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, sz, sp_b, sp_a in [('Heading 1', 18, 20, 10),
                             ('Heading 2', 14, 12, 6),
                             ('Heading 3', 12, 8, 4)]:
    s = doc.styles[name]
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    force_tnr(s)
    s.paragraph_format.space_before = Pt(sp_b)
    s.paragraph_format.space_after = Pt(sp_a)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing = 1.2
    s.paragraph_format.first_line_indent = Inches(0)

cs = doc.styles['Caption']
cs.font.size = Pt(10)
cs.font.color.rgb = RGBColor(0, 0, 0)
force_tnr(cs)
cs.paragraph_format.line_spacing = 1.0


# ===========================================================================
# Helper functions (shared by all part files)
# ===========================================================================

# --- Heuristic: which single character after `_` or `^` is treated as math?
# We avoid grabbing word characters like "_INR_" or "^USD" that appear in
# code-like tokens. The rule: a single-char sub/superscript is accepted only
# when the character is a digit, +, -, =, or a math-like letter (i, j, k, m,
# n, p, q, r, s, t, T) that actually means an index in our text. For longer
# tokens we always require braces: _{INR} or ^{+}.

_MATH_SINGLE_CHARS = set("0123456789+-=") | set("ijkmnpqrstT")


def _emit_runs(paragraph, text, *, font_size, font_name='Times New Roman',
               bold=False, italic=False):
    """Walk `text` and emit runs on `paragraph`. Any occurrence of `_` or `^`
    followed either by `{...}` or by a single character from `_MATH_SINGLE_CHARS`
    becomes a real Word subscript or superscript run. All other `_` or `^`
    characters are written literally. This lets us keep inline math tokens
    like Δy_t, CPT^{+}, N^{1/3}, oil^{INR}_{t} in plain strings without
    needing Unicode replacement."""
    if text is None:
        return
    i = 0
    n = len(text)
    buf = []

    def flush():
        if not buf:
            return
        chunk = ''.join(buf)
        r = paragraph.add_run(chunk)
        r.font.name = font_name
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rFonts.set(qn(k), font_name)
        r.font.size = Pt(font_size)
        r.bold = bold
        r.italic = italic
        buf.clear()

    def add_script(content, subscript):
        r = paragraph.add_run(content)
        r.font.name = font_name
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rFonts.set(qn(k), font_name)
        r.font.size = Pt(font_size)
        r.bold = bold
        r.italic = italic
        r.font.subscript = subscript
        r.font.superscript = not subscript

    while i < n:
        ch = text[i]
        if ch in ('_', '^') and i + 1 < n:
            nxt = text[i + 1]
            subscript = (ch == '_')
            if nxt == '{':
                j = text.find('}', i + 1)
                if j == -1:
                    buf.append(ch)
                    i += 1
                    continue
                content = text[i + 2:j]
                flush()
                add_script(content, subscript)
                i = j + 1
                continue
            if nxt in _MATH_SINGLE_CHARS:
                # Only treat as sub/superscript if the math char actually
                # terminates a token (next char is not alphanumeric).
                # This protects identifiers like D_petrol, oil_inr_pos, etc.
                after = text[i + 2] if (i + 2) < n else ''
                if not after.isalnum():
                    flush()
                    add_script(nxt, subscript)
                    i += 2
                    continue
        buf.append(ch)
        i += 1
    flush()


def add_para(text, indent=True, justify=True, size=11.5, bold=False, italic=False,
             align=None, space_after=None):
    p = doc.add_paragraph()
    if not indent:
        p.paragraph_format.first_line_indent = Inches(0)
    if align is not None:
        p.alignment = align
    elif not justify:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    _emit_runs(p, text, font_size=size, bold=bold, italic=italic)
    return p


def add_para_math(text, indent=True, justify=True, size=11.5, italic_default=False):
    """Alias retained for older call sites. Behaviour is identical to add_para
    now that add_para understands the sub/super tokens."""
    return add_para(text, indent=indent, justify=justify, size=size,
                    italic=italic_default)


def heading(text, level=1):
    style_name = f'Heading {level}'
    p = doc.add_paragraph(style=style_name)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'


def h1(text):
    heading(text, 1)


def h2(text):
    heading(text, 2)


def h3(text):
    heading(text, 3)


def bullets(items, size=11.5, space_after=2):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        _emit_runs(p, it, font_size=size)


def page_break():
    if doc.paragraphs:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    else:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_figure(path, caption, width_in=5.6):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Inches(width_in))
        cp = doc.add_paragraph(style='Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.first_line_indent = Inches(0)
        _emit_runs(cp, caption, font_size=10, italic=True)
    else:
        add_para(f"[FIGURE NOT FOUND: {path}]")


def add_table(headers, rows, caption=None, col_widths=None, font_size=10, note=None):
    if caption:
        cp = doc.add_paragraph(style='Caption')
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.first_line_indent = Inches(0)
        _emit_runs(cp, caption, font_size=10.5, bold=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        t.style = 'Table Grid'
    except KeyError:
        pass
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    if col_widths:
        usable_in = (doc.sections[0].page_width - doc.sections[0].left_margin
                     - doc.sections[0].right_margin) / 914400
        cw = list(col_widths)
        if sum(cw) > usable_in:
            scale = usable_in / sum(cw)
            cw = [w * scale for w in cw]
        dxa = [int(round(w * 1440)) for w in cw]
        old_grid = tbl.find(qn('w:tblGrid'))
        if old_grid is not None:
            tbl.remove(old_grid)
        grid = OxmlElement('w:tblGrid')
        for w in dxa:
            gc = OxmlElement('w:gridCol')
            gc.set(qn('w:w'), str(w))
            grid.append(gc)
        tbl.insert(list(tbl).index(tblPr) + 1, grid)
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW = OxmlElement('w:tcW')
                    tcPr.append(tcW)
                tcW.set(qn('w:w'), str(dxa[ci]))
                tcW.set(qn('w:type'), 'dxa')

    def write_cell(cell, text, bold=False, align='center'):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, v in (('top', 60), ('start', 80), ('bottom', 60), ('end', 80)):
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(v))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)
        v_align = OxmlElement('w:vAlign')
        v_align.set(qn('w:val'), 'center')
        tcPr.append(v_align)
        cell.text = ''
        para = cell.paragraphs[0]
        para.paragraph_format.first_line_indent = Inches(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        para.alignment = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        _emit_runs(para, str(text), font_size=font_size, bold=bold)

    for i, h in enumerate(headers):
        write_cell(t.rows[0].cells[i], h, bold=True, align='center')
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            write_cell(t.rows[ri].cells[ci], str(val),
                       align='left' if ci == 0 else 'center')
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))
    if note:
        np_ = doc.add_paragraph()
        np_.paragraph_format.first_line_indent = Inches(0)
        np_.paragraph_format.space_before = Pt(2)
        np_.paragraph_format.space_after = Pt(8)
        np_.paragraph_format.line_spacing = 1.0
        _emit_runs(np_, "Note: " + note, font_size=9, italic=True)


# --- OMML helpers (same style as the main builders) -------------------------
M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def add_equation(omml_inner, label=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p_elem = p._p
    pPr = p_elem.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for val, pos in [('center', '4230'), ('right', '8460')]:
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), val)
        tab.set(qn('w:pos'), pos)
        tabs.append(tab)
    pPr.append(tabs)
    p_elem.append(parse_xml(f'<w:r xmlns:w="{W_NS}"><w:tab/></w:r>'))
    p_elem.append(parse_xml(f'<m:oMath xmlns:m="{M_NS}">{omml_inner}</m:oMath>'))
    if label:
        p_elem.append(parse_xml(
            f'<w:r xmlns:w="{W_NS}"><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
            f'<w:sz w:val="22"/></w:rPr><w:tab/><w:t xml:space="preserve">{label}</w:t></w:r>'))
    return p


def mr(t, italic=True):
    t = t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if not italic else ''
    return f'<m:r>{sty}<m:t xml:space="preserve">{t}</m:t></m:r>'


def msub(b, s):
    return f'<m:sSub><m:e>{b}</m:e><m:sub>{s}</m:sub></m:sSub>'


def msup(b, s):
    return f'<m:sSup><m:e>{b}</m:e><m:sup>{s}</m:sup></m:sSup>'


def msubsup(b, sb, sp):
    return f'<m:sSubSup><m:e>{b}</m:e><m:sub>{sb}</m:sub><m:sup>{sp}</m:sup></m:sSubSup>'


def mfenced(i):
    return f'<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{i}</m:e></m:d>'


def mnary_sum(sub, sup, body):
    return ('<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/>'
            '<m:subHide m:val="0"/><m:supHide m:val="0"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{sup}</m:sup><m:e>{body}</m:e></m:nary>')


def add_code_block(text, size=9.5):
    """Monospaced pseudo-code paragraph with light gray shading."""
    lines = text.rstrip("\n").split("\n")
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        # light gray shading
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F4F4F4')
        pPr.append(shd)
        r = p.add_run(line if line else " ")
        r.font.name = 'Consolas'
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rFonts.set(qn(k), 'Consolas')
        r.font.size = Pt(size)
    # trailing spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)


def add_callout(title, text):
    """A light gray paragraph with a bold title and italic body, to highlight tips."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'EEF2FA')
    pPr.append(shd)
    _emit_runs(p, title + ": ", font_size=11, bold=True)
    _emit_runs(p, text, font_size=11, italic=True)


# ===========================================================================
# Title page
# ===========================================================================

def title_page():
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.first_line_indent = Inches(0)
    tp.paragraph_format.space_before = Pt(90)
    tp.paragraph_format.space_after = Pt(18)
    tr = tp.add_run("Study Guide")
    tr.bold = True
    tr.font.name = 'Times New Roman'
    tr.font.size = Pt(28)

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.first_line_indent = Inches(0)
    sp.paragraph_format.space_after = Pt(8)
    sr = sp.add_run("From wholesale prices to consumer inflation: "
                    "layered pass-through of global oil shocks in India")
    sr.bold = True
    sr.font.name = 'Times New Roman'
    sr.font.size = Pt(16)

    ap = doc.add_paragraph()
    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ap.paragraph_format.first_line_indent = Inches(0)
    ap.paragraph_format.space_after = Pt(6)
    ar = ap.add_run("A plain-language walk-through for presentation and viva")
    ar.italic = True
    ar.font.name = 'Times New Roman'
    ar.font.size = Pt(13)

    au = doc.add_paragraph()
    au.alignment = WD_ALIGN_PARAGRAPH.CENTER
    au.paragraph_format.first_line_indent = Inches(0)
    au.paragraph_format.space_before = Pt(30)
    au.paragraph_format.space_after = Pt(6)
    aur = au.add_run("Companion to the dissertation and journal paper by Aniket Pandey")
    aur.font.name = 'Times New Roman'
    aur.font.size = Pt(12)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = Inches(0)
    note.paragraph_format.space_before = Pt(36)
    nr = note.add_run(
        "Read this side by side with aniket-dissertation.docx and journal-paper.docx. "
        "Every result cited here comes from the model outputs in models/cpi and models/wpi."
    )
    nr.italic = True
    nr.font.name = 'Times New Roman'
    nr.font.size = Pt(11)

    page_break()


def toc():
    h1("How to use this guide")
    add_para(
        "This guide is written for you, the student, and for anyone grading "
        "or listening to your defence. It is not a second dissertation. It is a "
        "companion. Read it in order the first time through. After that, jump "
        "to whichever section a specific question points to. Every Greek "
        "letter, every test, every model choice is explained in plain words "
        "before any formula is used."
    )
    bullets([
        "Part 1: The big picture in one page.",
        "Part 2: Background. India, oil, WPI, CPI, and why the dollar matters.",
        "Part 3: Research question, objectives, and contribution.",
        "Part 4: Data sources, full forms, and what each series is.",
        "Part 5: Variables. Log differences, shock splits, and the rupee oil price.",
        "Part 6: Methodology. ADL models, OLS, and what α, β, γ, μ, ε, φ mean.",
        "Part 7: Inference. HAC standard errors, bootstrap, Wald, Granger.",
        "Part 8: Diagnostics. Unit roots, breaks, BG, RESET, CUSUM, the model gate.",
        "Part 9: Results. Layer by layer with every number spelled out.",
        "Part 10: Robustness and limitations.",
        "Part 11: Conclusion and what the teacher will ask.",
        "Part 12: Viva question bank with model answers.",
        "Part 13: Glossary and cheat sheet.",
    ])
    add_callout(
        "Tip",
        "If a single question lands on one topic, turn to the Viva question bank in "
        "Part 12 first, then read the relevant methodology or results section for "
        "supporting detail."
    )
    page_break()


def main():
    title_page()
    toc()

    import _guide_part1_overview as p1
    import _guide_part2_background as p2
    import _guide_part3_research_design as p3
    import _guide_part4_data as p4
    import _guide_part5_variables as p5
    import _guide_part6_methodology as p6
    import _guide_part7_inference as p7
    import _guide_part8_diagnostics as p8
    import _guide_part9_results as p9
    import _guide_part10_robustness as p10
    import _guide_part11_conclusion as p11
    import _guide_part12_qa as p12
    import _guide_part13_glossary as p13

    helpers = dict(
        doc=doc,
        add_para=add_para,
        add_para_math=add_para_math,
        h1=h1, h2=h2, h3=h3,
        bullets=bullets,
        page_break=page_break,
        add_figure=add_figure,
        add_table=add_table,
        add_equation=add_equation,
        add_code_block=add_code_block,
        add_callout=add_callout,
        emit_runs=_emit_runs,
        mr=mr, msub=msub, msup=msup, msubsup=msubsup,
        mfenced=mfenced, mnary_sum=mnary_sum,
        ROOT=str(ROOT),
    )

    for part in (p1, p2, p3, p4, p5, p6, p7, p8, p9, p10, p11, p12, p13):
        part.build(helpers)

    doc.save(OUT_DOCX)
    print(f"Saved: {OUT_DOCX}")

    # Convert to PDF via soffice if available
    soffice = "/opt/homebrew/bin/soffice"
    if os.path.exists(soffice):
        try:
            subprocess.run([
                soffice, "--headless", "--convert-to", "pdf",
                "--outdir", str(ROOT), str(OUT_DOCX),
            ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            print(f"Saved: {OUT_PDF}")
        except subprocess.CalledProcessError as exc:
            print(f"soffice conversion failed: {exc}")
    else:
        print("soffice not found; skipped PDF conversion.")


if __name__ == "__main__":
    main()
